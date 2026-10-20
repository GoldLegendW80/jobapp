"""The styles of a Word document, for the design capture of an uploaded CV (task D25a of the
documents plan; SPEC-documents.md, "Design capture of an upload", step 1). Runs only inside the
locked down child of the parser service: main.py imports this file after the lockdown and calls
`handle`. The bytes are hostile, and so is every name and every attribute found in them.

    handle(data, options, budget) -> the whole answer, as app/ingestion/parser_client.py reads it

docx.py reads the text and deliberately inflates nothing else; this reads what it leaves: the
theme (`word/theme/theme1.xml`: its colours and its two fonts), the styles (`word/styles.xml`:
the defaults and each style's run and paragraph properties) and, of `word/document.xml`, the
layout and the direct formatting, never a word of the text: the sections (page size, margins,
columns), the tables (how many, the most columns of one), the drawings, how many paragraphs use
each style, and how many characters each combination of direct run properties formats. The
archive is opened under the caps of zipcaps.py, the macros are refused as docx.py refuses them,
and the XML goes through docx.py's own parser and its counts of names and depth.

    {"ok": true, "kind": "design_docx", "truncated": bool, "notes": [{"code": ...}],
     "theme": {"major": font | null, "minor": font | null, "colours": {"accent1": "#rrggbb"}},
     "defaults": PROPS,
     "styles": [{"id", "name", "type", "based_on", "default": bool, "outline": 0..9 | null,
                 **PROPS}],
     "sections": [{"width", "height", "margins": [top, right, bottom, left], "columns"}],
     "tables": N, "table_columns": N, "drawings": N, "paragraphs": N,
     "paragraph_styles": [{"style": id, "count": N}],
     "formats": [{"style": id | null, "font", "size", "bold", "italic", "caps", "colour",
                  "chars": N}]}

PROPS are `font`, `size`, `bold`, `italic`, `caps`, `colour`, `before`, `after`, `line`,
`line_rule`, `align`, `border` and `shading`, each null when the part does not set it. Lengths
are integers in tenths of a point (Word's half points and twentieths converted); `line` is in
hundredths of a line when `line_rule` is "auto", else in tenths of a point. A theme font or a
theme colour is resolved through the theme. A DOCX has no pages: nothing here is a box.

Caps: STYLES_MAX styles, SECTIONS_MAX sections, KEYS_MAX distinct style uses and formats
counted, the SHOWN most used of each returned. Each one met is a note and `truncated`. A count
is given as COUNT_MAX at most, and a value the client would refuse is left out (null), so one
odd value never costs the whole answer.

What comes back otherwise is docx.py's: `not_docx`, `macros`, `unreadable`, `bomb`,
`nested_archive`, `protected`, `too_slow`.

Only main.py may import this file in production. The tests import it too, to call `handle` in
their own process on files they make.
"""

import docx  # sandbox/docx.py, beside this file: its parser, its counts, its refusals
import zipcaps
from defusedxml import ElementTree
from defusedxml.common import DefusedXmlException

THEME = "word/theme/theme1.xml"
STYLES = "word/styles.xml"
STYLES_MAX = 400
SECTIONS_MAX = 20
KEYS_MAX = 2_000
SHOWN = 40
NAME_MAX = 64
TWIPS_MAX = 200_000  # a length Word may write, in twentieths of a point (139 inches)
COUNT_MAX = 10_000_000  # parser_client.DESIGN_COUNT: a count past it is given as this

_A = frozenset(
    {
        "http://schemas.openxmlformats.org/drawingml/2006/main",
        "http://purl.oclc.org/ooxml/drawingml/main",
    }
)
_SCHEME = frozenset(
    (
        "dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4", "accent5",
        "accent6", "hlink", "folHlink",
    )
)  # fmt: skip
# w:themeColor -> the colour of the theme's scheme it names.
_THEME_COLOURS = {
    "dark1": "dk1", "text1": "dk1", "light1": "lt1", "background1": "lt1", "dark2": "dk2",
    "text2": "dk2", "light2": "lt2", "background2": "lt2", "accent1": "accent1",
    "accent2": "accent2", "accent3": "accent3", "accent4": "accent4", "accent5": "accent5",
    "accent6": "accent6", "hyperlink": "hlink", "followedHyperlink": "folHlink",
}  # fmt: skip
_ALIGN = {
    "left": "left", "start": "left", "center": "center", "right": "right", "end": "right",
    "both": "justify", "distribute": "justify",
}  # fmt: skip
_LINE_RULES = {"auto": "auto", "exact": "exact", "atLeast": "at_least"}
_STYLE_TYPES = frozenset({"paragraph", "character", "table", "numbering"})
_FALSE = frozenset({"0", "false", "off", "none"})
_HEX_DIGITS = frozenset("0123456789abcdefABCDEF")
PROPS = (
    "font", "size", "bold", "italic", "caps", "colour", "before", "after", "line", "line_rule",
    "align", "border", "shading",
)  # fmt: skip
RUN_PROPS = ("font", "size", "bold", "italic", "caps", "colour")


def _props() -> dict:
    return dict.fromkeys(PROPS)


def _name(value: str | None) -> str | None:
    if value is None:
        return None
    kept = "".join(ch for ch in value if ch.isprintable() and not 0xD800 <= ord(ch) <= 0xDFFF)
    kept = kept.strip()[:NAME_MAX]
    return kept or None


def _number(value: str | None, high: int = TWIPS_MAX) -> int | None:
    """The whole number `value` writes in ASCII digits, with one minus sign at most, when it
    lies in 0..high; None for anything else ("²", "--5", "12pt", "1e3", "٣")."""
    if value is None or len(value) > 9:
        return None
    text = value.strip()
    digits = text.removeprefix("-")
    if not (digits.isascii() and digits.isdigit()):
        return None
    number = int(text)
    return number if 0 <= number <= high else None


def _hex(value: str | None) -> str | None:
    """`#rrggbb` when `value` is six ASCII hex digits and nothing else; None for anything else
    ("+abcde", "ab_cde", " abcd ", "0xabcd", "١٢٣٤٥٦")."""
    if value is None or len(value) != 6 or not set(value) <= _HEX_DIGITS:
        return None
    return "#" + value.lower()


def _counted(number: int) -> int:
    return min(number, COUNT_MAX)


def _local(key: str, spaces: frozenset[str]) -> str | None:
    """The local name of a tag or attribute in one of `spaces`, None otherwise."""
    space, name = docx._split(key)
    return name if space in spaces else None


def _attribute(attributes: dict[str, str], name: str) -> str | None:
    for key, value in attributes.items():
        space, local = docx._split(key)
        if local == name and (space in docx._W or not space):
            return value
    return None


class _Target(docx._Bounded):
    """What the three targets share: docx.py's counts, a stack of the open elements' names."""

    def __init__(self) -> None:
        super().__init__()
        self.stack: list[str] = []

    def start(self, tag: str, attributes: dict[str, str]) -> None:
        self.opened(tag, attributes)
        space, name = docx._split(tag)
        local = name if space in docx._W else (f"a:{name}" if space in _A else f"?:{name}")
        self.stack.append(local)
        self.on_open(local, attributes)

    def end(self, tag: str) -> None:
        self.depth -= 1
        self.on_close(self.stack.pop() if self.stack else "")

    def data(self, text: str) -> None:
        self.on_text(text)

    def close(self) -> None:
        return None

    def on_open(self, name: str, attributes: dict[str, str]) -> None:
        return None

    def on_close(self, name: str) -> None:
        return None

    def on_text(self, text: str) -> None:
        return None


class _Theme(_Target):
    def __init__(self) -> None:
        super().__init__()
        self.colours: dict[str, str] = {}
        self.fonts: dict[str, str | None] = {"major": None, "minor": None}
        self.slot: str | None = None

    def on_open(self, name: str, attributes: dict[str, str]) -> None:
        if not name.startswith("a:"):
            return
        name = name[2:]
        if name in _SCHEME and "a:clrScheme" in self.stack:
            self.slot = name
        elif name in ("srgbClr", "sysClr") and self.slot and self.slot not in self.colours:
            value = attributes.get("val") if name == "srgbClr" else attributes.get("lastClr")
            colour = _hex(value)
            if colour:
                self.colours[self.slot] = colour
        elif name == "latin" and "a:fontScheme" in self.stack:
            which = "major" if "a:majorFont" in self.stack else "minor"
            if "a:minorFont" in self.stack or "a:majorFont" in self.stack:
                self.fonts[which] = self.fonts[which] or _name(attributes.get("typeface"))

    def on_close(self, name: str) -> None:
        if name[2:] == self.slot:
            self.slot = None


class _Properties:
    """Reads the run and paragraph properties of one element into a PROPS dictionary."""

    def __init__(self, theme: _Theme | None) -> None:
        self.theme = theme

    def font(self, attributes: dict[str, str]) -> str | None:
        for key in ("ascii", "hAnsi"):
            found = _name(_attribute(attributes, key))
            if found:
                return found
        for key in ("asciiTheme", "hAnsiTheme"):
            found = _attribute(attributes, key) or ""
            if self.theme is not None and found.startswith(("major", "minor")):
                return self.theme.fonts["major" if found.startswith("major") else "minor"]
        return None

    def colour(self, attributes: dict[str, str], key: str = "val") -> str | None:
        found = _hex(_attribute(attributes, key))
        if found:
            return found
        slot = _THEME_COLOURS.get(_attribute(attributes, "themeColor") or "")
        if slot and self.theme is not None:
            return self.theme.colours.get(slot)
        return None

    def read(self, props: dict, name: str, attributes: dict[str, str], within: str) -> None:
        """Sets in `props` what the element `name`, a child of `within`, says."""
        value = _attribute(attributes, "val")
        if within == "rPr":
            if name == "rFonts":
                props["font"] = self.font(attributes) or props["font"]
            elif name == "sz":
                half = _number(value, 3_276)
                props["size"] = half * 5 if half else props["size"]
            elif name in ("b", "i", "caps", "smallCaps"):
                key = {"b": "bold", "i": "italic"}.get(name, "caps")
                props[key] = (value or "").lower() not in _FALSE
            elif name == "color":
                props["colour"] = self.colour(attributes)
        elif within == "pPr":
            if name == "spacing":
                for key in ("before", "after"):
                    twips = _number(_attribute(attributes, key))
                    props[key] = twips // 2 if twips is not None else props[key]
                line = _number(_attribute(attributes, "line"))
                rule = _LINE_RULES.get(_attribute(attributes, "lineRule") or "auto")
                if line is not None and rule:
                    props["line_rule"] = rule
                    props["line"] = line * 100 // 240 if rule == "auto" else line // 2
            elif name == "jc":
                props["align"] = _ALIGN.get(value or "", props["align"])
            elif name == "shd":
                props["shading"] = _hex(_attribute(attributes, "fill")) or props["shading"]
        elif within == "pBdr" and name == "bottom" and (value or "none") not in ("none", "nil"):
            props["border"] = self.colour(attributes, "color") or "#000000"


class _Styles(_Target):
    def __init__(self, theme: _Theme | None) -> None:
        super().__init__()
        self.reader = _Properties(theme)
        self.defaults = _props()
        self.styles: list[dict] = []
        self.current: dict | None = None
        self.full = False

    def _props_here(self) -> tuple[dict | None, str]:
        """The PROPS an element at the top of the stack writes to, and its parent's name."""
        path = self.stack
        if len(path) < 3:
            return None, ""
        within = path[-2]
        if "tblStylePr" in path:
            return None, ""  # the conditional formatting of a table style: not the style
        if "docDefaults" in path:
            return self.defaults, within
        if self.current is not None and path[-3:-1] in (["style", "rPr"], ["style", "pPr"]):
            return self.current, within
        if self.current is not None and path[-4:-1] == ["style", "pPr", "pBdr"]:
            return self.current, within
        return None, ""

    def on_open(self, name: str, attributes: dict[str, str]) -> None:
        if name == "style" and self.stack[:-1] == ["styles"]:
            if len(self.styles) == STYLES_MAX:
                self.full = True
                self.current = None
                return
            kind = _attribute(attributes, "type")
            self.current = {
                "id": _name(_attribute(attributes, "styleId")) or "",
                "name": None,
                "type": kind if kind in _STYLE_TYPES else "paragraph",
                "based_on": None,
                "default": (_attribute(attributes, "default") or "").lower() in ("1", "true", "on"),
                "outline": None,
                **_props(),
            }
            return
        if self.current is not None and self.stack[-2:-1] == ["style"]:
            value = _attribute(attributes, "val")
            if name == "name":
                self.current["name"] = _name(value)
            elif name == "basedOn":
                self.current["based_on"] = _name(value)
        if (
            self.current is not None
            and name == "outlineLvl"
            and self.stack[-3:-1]
            == [
                "style",
                "pPr",
            ]
        ):
            level = _number(_attribute(attributes, "val"), 9)
            self.current["outline"] = level
        props, within = self._props_here()
        if props is not None:
            self.reader.read(props, name, attributes, within)

    def on_close(self, name: str) -> None:
        if name == "style" and self.current is not None and self.stack == ["styles"]:
            if self.current["id"]:
                self.styles.append(self.current)
            self.current = None


class _Document(_Target):
    def __init__(self, theme: _Theme | None) -> None:
        super().__init__()
        self.reader = _Properties(theme)
        self.sections: list[dict] = []
        self.section: dict | None = None
        self.tables = self.table_columns = self.drawings = self.paragraphs = 0
        self.columns = 0
        self.paragraph_styles: dict[str, int] = {}
        self.formats: dict[tuple, int] = {}
        self.run: dict | None = None
        self.run_style: str | None = None
        self.key: tuple | None = None
        self.notes: set[str] = set()
        self.skip = 0  # depth of an mc:Fallback being skipped

    def _count(self, table: dict, key, amount: int) -> None:
        if key in table:
            table[key] = _counted(table[key] + amount)
        elif len(table) < KEYS_MAX:
            table[key] = _counted(amount)
        else:
            self.notes.add("formats_limit")

    def on_open(self, name: str, attributes: dict[str, str]) -> None:
        if self.skip:
            return
        if name == "?:Fallback":
            self.skip = len(self.stack)
            return
        parent = self.stack[-2] if len(self.stack) > 1 else ""
        if name == "p":
            self.paragraphs += 1
        elif name == "pStyle" and parent == "pPr":
            self._count(self.paragraph_styles, _name(_attribute(attributes, "val")) or "", 1)
        elif name == "tbl":
            self.tables += 1
        elif name == "tblGrid":
            self.columns = 0
        elif name == "gridCol":
            self.columns += 1
            self.table_columns = max(self.table_columns, self.columns)
        elif name in ("drawing", "pict"):
            self.drawings += 1
        elif name == "r":
            self.run, self.run_style, self.key = _props(), None, None
        elif name == "rStyle" and self.run is not None:
            self.run_style = _name(_attribute(attributes, "val"))
        elif self.run is not None and parent == "rPr" and "r" in self.stack[-3:-2]:
            self.reader.read(self.run, name, attributes, "rPr")
        elif name == "sectPr":
            if len(self.sections) < SECTIONS_MAX:
                self.section = {"width": None, "height": None, "margins": None, "columns": 1}
            else:
                self.notes.add("sections_limit")
        elif self.section is not None and parent == "sectPr":
            self._section(name, attributes)

    def _section(self, name: str, attributes: dict[str, str]) -> None:
        section = self.section
        if name == "pgSz":
            width, height = (
                _number(_attribute(attributes, "w")),
                _number(_attribute(attributes, "h")),
            )
            section["width"] = width // 2 if width else None
            section["height"] = height // 2 if height else None
        elif name == "pgMar":
            sides = [
                _number(_attribute(attributes, side)) for side in ("top", "right", "bottom", "left")
            ]
            if all(side is not None for side in sides):
                section["margins"] = [side // 2 for side in sides]
        elif name == "cols":
            count = _number(_attribute(attributes, "num"), 45)
            section["columns"] = count if count else 1

    def on_close(self, name: str) -> None:
        if self.skip:
            if len(self.stack) + 1 == self.skip:
                self.skip = 0
            return
        if name == "sectPr" and self.section is not None:
            self.sections.append(self.section)
            self.section = None
        elif name == "rPr" and self.run is not None:
            self.key = (self.run_style, *(self.run[key] for key in RUN_PROPS))
        elif name == "r":
            self.run = None

    def on_text(self, text: str) -> None:
        if self.skip or not self.stack or self.stack[-1] != "t" or self.run is None:
            return
        key = self.key or (self.run_style, *(None for _ in RUN_PROPS))
        self._count(self.formats, key, len(text))


def _parse(archive: zipcaps.Archive, part: str, target, budget) -> bool | None:
    """True when `part` was read whole, False when the budget ended it, None when it is
    missing or is not XML."""
    if part not in archive.names:
        return None
    try:
        return docx.parse(archive.read(part), target, budget)
    except ElementTree.ParseError:
        return None


def _shown(counts: dict, limit: int = SHOWN) -> list:
    return sorted(counts.items(), key=lambda item: (-item[1], str(item[0])))[:limit]


def read(archive: zipcaps.Archive, budget) -> dict:
    notes: set[str] = set()
    theme = _Theme()
    if _parse(archive, THEME, theme, budget) is False:
        notes.add("time_limit")
    styles = _Styles(theme)
    if "time_limit" not in notes and _parse(archive, STYLES, styles, budget) is False:
        notes.add("time_limit")
    if styles.full:
        notes.add("styles_limit")
    document = _Document(theme)
    if "time_limit" not in notes:
        whole = _parse(archive, docx.DOCUMENT, document, budget)
        if whole is None:
            return docx.failure("unreadable")
        if whole is False:
            notes.add("time_limit")
    if "time_limit" in notes and not styles.styles and not document.paragraphs:
        return docx.failure("too_slow")
    notes |= document.notes
    has_theme = bool(theme.colours) or any(theme.fonts.values())
    return {
        "ok": True,
        "kind": "design_docx",
        "truncated": bool(notes),
        "notes": [{"code": code} for code in sorted(notes)],
        "theme": {**theme.fonts, "colours": theme.colours} if has_theme else None,
        "defaults": styles.defaults,
        "styles": styles.styles,
        "sections": document.sections,
        "tables": _counted(document.tables),
        "table_columns": _counted(document.table_columns),
        "drawings": _counted(document.drawings),
        "paragraphs": _counted(document.paragraphs),
        "paragraph_styles": [
            {"style": style, "count": count} for style, count in _shown(document.paragraph_styles)
        ],
        "formats": [
            {"style": key[0], **dict(zip(RUN_PROPS, key[1:], strict=True)), "chars": chars}
            for key, chars in _shown(document.formats)
        ],
    }


def handle(data: bytes, options: dict, budget) -> dict:
    """The styles and the layout of `data`, a Word document. `options` must be empty."""
    if options:
        return {"ok": False, "failure": "server", "code": "handler_error"}
    try:
        archive = zipcaps.Archive(data)
        lowered = {name.lower() for name in archive.names}
        if any(name.rsplit("/", 1)[-1] == docx.MACROS_NAME for name in lowered):
            return docx.failure("macros")
        if docx.CONTENT_TYPES not in archive.names:
            return docx.failure("not_docx")
        if docx.DOCUMENT not in archive.names:
            damaged = any(name.startswith("word/") for name in lowered)
            return docx.failure("unreadable" if damaged else "not_docx")
        if docx.MACROS_TYPE in archive.read(docx.CONTENT_TYPES):
            return docx.failure("macros")
        return read(archive, budget)
    except zipcaps.Refused as refused:
        return docx.failure(refused.code)
    except (DefusedXmlException, MemoryError):
        return docx.failure("bomb")
